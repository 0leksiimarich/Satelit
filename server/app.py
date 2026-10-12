# ============================================================
# app.py — сервер «Сателіт» (FastAPI, Render.com)
# ------------------------------------------------------------
# Що вміє:
#  • вхід і реєстрація (email + пароль; код із пошти для реєстрації)
#  • каталог товарів і замовлення
#  • адмін-панель: акаунти, товари, імпорт таблиці (.xlsx / .csv),
#    перевірка рядків через ШІ (OpenRouter)
# Дані зберігаються у JSON-файлі (SATELIT_DATA_DIR). Щоб дані не
# зникали після перезапуску, потрібен постійний Disk на Render.
# ============================================================
import copy
import csv
import hashlib
import io
import json
import os
import re
import secrets
import smtplib
import ssl
import threading
import time
from datetime import datetime, timezone
from email.message import EmailMessage

import requests
from fastapi import FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from openpyxl import load_workbook
from pydantic import BaseModel, Field

app = FastAPI(title="Satelit Backend")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in os.environ.get("SATELIT_CORS_ORIGINS", "*").split(",") if o.strip()],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.exception_handler(RequestValidationError)
async def validation_error_handler(request, exc):
    return JSONResponse(
        status_code=422,
        content={"detail": "Перевірте, будь ласка, правильність заповнення полів."},
    )


# ---------- Налаштування (змінні середовища на Render) ----------
DATA_DIR = os.environ.get("SATELIT_DATA_DIR", "data")
DATA_FILE = os.environ.get("SATELIT_DATA_FILE", os.path.join(DATA_DIR, "satelit_data.json"))
SUPPORT_EMAIL = os.environ.get("SATELIT_SUPPORT_EMAIL", "satelit-magazin@ukr.net")
# Пошта власника: акаунт із цією адресою автоматично стає адміном (без командного рядка)
ADMIN_EMAIL = (os.environ.get("SATELIT_ADMIN_EMAIL", "") or "").strip().lower()

SMTP_HOST = os.environ.get("SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "465"))
SMTP_USER = os.environ.get("SMTP_USER", "")
SMTP_PASS = os.environ.get("SMTP_PASS", "")
SMTP_FROM = os.environ.get("SMTP_FROM", SMTP_USER or SUPPORT_EMAIL)

OPENROUTER_API_KEY = os.environ.get("OPENROUTER_API_KEY", "")
OPENROUTER_MODEL = os.environ.get("OPENROUTER_MODEL", "google/gemma-4-31b-it:free")

OTP_TTL = 15 * 60          # код дійсний 15 хвилин
OTP_RESEND_AFTER = 60      # повторно не раніше ніж через 60 с
MAX_OTP_ATTEMPTS = 5

CATEGORY_IDS = {"batteries", "tvbox", "antennas", "cables", "accessories", "other"}
CATEGORY_ALIASES = {
    "батарейки": "batteries", "батарейка": "batteries",
    "tv box": "tvbox", "тв бокс": "tvbox", "тв-бокс": "tvbox", "тv box": "tvbox",
    "антени": "antennas", "антена": "antennas",
    "кабелі": "cables", "кабель": "cables",
    "аксесуари": "accessories", "інше": "other",
}
STATUSES = ["Нове", "Підтверджено", "Готується", "Відправлено", "Виконано", "Скасовано"]
DELIVERY_METHODS = {"pickup", "novaposhta", "ukrposhta"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
DEFAULT_IMAGE = "assets/img-placeholder.svg"


# ============================================================
# Сховище (JSON-файл)
# ============================================================
DB_LOCK = threading.RLock()
DEFAULT_DATA = {"users": [], "products": [], "orders": [], "tokens": {}, "otps": []}


def _load():
    data = {}
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, encoding="utf-8") as f:
            data = json.load(f)
    for key, value in DEFAULT_DATA.items():
        data.setdefault(key, copy.deepcopy(value))
    return data


def db():
    with DB_LOCK:
        return _load()


def write_db(data):
    with DB_LOCK:
        folder = os.path.dirname(DATA_FILE)
        if folder:
            os.makedirs(folder, exist_ok=True)
        tmp = DATA_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        os.replace(tmp, DATA_FILE)


def now_iso():
    return datetime.now(timezone.utc).isoformat()


def norm_email(value):
    return (value or "").strip().lower()


# ============================================================
# Паролі, токени, користувачі
# ============================================================
def hash_password(pw):
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), bytes.fromhex(salt), 100_000).hex()
    return f"{salt}${digest}"


def verify_password(pw, stored):
    try:
        salt, digest = stored.split("$", 1)
    except (ValueError, AttributeError):
        return False
    check = hashlib.pbkdf2_hmac("sha256", pw.encode("utf-8"), bytes.fromhex(salt), 100_000).hex()
    return secrets.compare_digest(check, digest)


def make_token(data, user_id):
    token = secrets.token_urlsafe(32)
    data["tokens"][token] = user_id
    return token


def user_from_request(request, data):
    auth = request.headers.get("authorization", "")
    if not auth.lower().startswith("bearer "):
        return None
    uid = data["tokens"].get(auth[7:].strip())
    if not uid:
        return None
    return next((u for u in data["users"] if u["id"] == uid), None)


def require_user(request, data):
    user = user_from_request(request, data)
    if not user:
        raise HTTPException(401, "Потрібно увійти в акаунт")
    if user.get("blocked"):
        raise HTTPException(403, "Акаунт заблоковано. Зверніться до підтримки.")
    return user


def require_admin(request, data):
    user = require_user(request, data)
    if not user.get("isAdmin"):
        raise HTTPException(403, "Ця дія доступна лише адміністраторам")
    return user


def public_user(u):
    return {
        "id": u["id"],
        "email": u["email"],
        "displayName": u.get("name", ""),
        "phone": u.get("phone", ""),
        "isAdmin": bool(u.get("isAdmin")),
        "blocked": bool(u.get("blocked")),
        "createdAt": u.get("createdAt", ""),
    }


# ============================================================
# Коди з пошти (OTP) і відправка листів
# ============================================================
def _otp_hash(email, purpose, code):
    return hashlib.sha256(f"{email}|{purpose}|{code}".encode("utf-8")).hexdigest()


def otp_issue(data, email, purpose):
    now = time.time()
    for o in data["otps"]:
        if o["email"] == email and o["purpose"] == purpose and now - o["created"] < OTP_RESEND_AFTER:
            raise HTTPException(429, "Код уже надіслано. Зачекайте хвилину і спробуйте знову.")
    data["otps"] = [
        o for o in data["otps"]
        if o["expires"] > now and not (o["email"] == email and o["purpose"] == purpose)
    ]
    code = f"{secrets.randbelow(1_000_000):06d}"
    data["otps"].append({
        "email": email, "purpose": purpose,
        "hash": _otp_hash(email, purpose, code),
        "created": now, "expires": now + OTP_TTL, "attempts": 0,
    })
    return code


def otp_check(data, email, purpose, code):
    """Повертає текст помилки або None, якщо код вірний. Викликач має зберегти data."""
    now = time.time()
    entry = next((o for o in data["otps"] if o["email"] == email and o["purpose"] == purpose), None)
    if not entry:
        return "Спершу запросіть код на пошту"
    if entry["expires"] < now:
        data["otps"].remove(entry)
        return "Код застарів. Запросіть новий."
    if entry["attempts"] >= MAX_OTP_ATTEMPTS:
        data["otps"].remove(entry)
        return "Забагато спроб. Запросіть новий код."
    if not secrets.compare_digest(entry["hash"], _otp_hash(email, purpose, code)):
        entry["attempts"] += 1
        return "Невірний код"
    data["otps"].remove(entry)
    return None


def send_email(to, subject, text):
    """True — лист надіслано. False — SMTP не налаштовано (демо-режим)."""
    if not (SMTP_HOST and SMTP_USER and SMTP_PASS):
        return False
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = SMTP_FROM
    msg["To"] = to
    msg.set_content(text)
    try:
        ctx = ssl.create_default_context()
        if SMTP_PORT == 465:
            with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, context=ctx, timeout=20) as s:
                s.login(SMTP_USER, SMTP_PASS)
                s.send_message(msg)
        else:
            with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as s:
                s.starttls(context=ctx)
                s.login(SMTP_USER, SMTP_PASS)
                s.send_message(msg)
    except (smtplib.SMTPException, OSError) as e:
        print("SMTP error:", e)
        raise HTTPException(502, "Не вдалося надіслати лист із кодом. Спробуйте пізніше.")
    return True


# ============================================================
# Моделі запитів
# ============================================================
class OtpSendIn(BaseModel):
    email: str
    purpose: str = "register"


class RegisterIn(BaseModel):
    email: str
    password: str
    name: str
    phone: str = ""
    code: str = ""


class LoginIn(BaseModel):
    email: str
    password: str
    code: str = ""


class ProductIn(BaseModel):
    name: str
    description: str = ""
    price: float = Field(ge=0)
    category: str = "other"
    image: str = DEFAULT_IMAGE
    stock: int = Field(default=0, ge=0)
    featured: bool = False
    sale: bool = False


class OrderItemIn(BaseModel):
    id: str
    qty: int = Field(default=1, ge=1, le=999)


class OrderIn(BaseModel):
    items: list[OrderItemIn]
    customer_name: str = ""
    phone: str = ""
    email: str = ""
    delivery_method: str = "pickup"
    address: str = ""


class StatusIn(BaseModel):
    status: str


class RoleIn(BaseModel):
    isAdmin: bool


class BlockIn(BaseModel):
    blocked: bool


class PasswordIn(BaseModel):
    password: str


# ============================================================
# Загальне
# ============================================================
@app.get("/api/health")
def health():
    return {"ok": True, "time": now_iso()}


# ============================================================
# Авторизація
# ============================================================
@app.post("/api/auth/otp/send")
def otp_send(body: OtpSendIn):
    email = norm_email(body.email)
    purpose = body.purpose
    if purpose not in ("register", "login"):
        raise HTTPException(400, "Невірна мета коду")
    if not EMAIL_RE.match(email):
        raise HTTPException(400, "Вкажіть коректну електронну пошту")
    with DB_LOCK:
        data = _load()
        exists = any(u["email"] == email for u in data["users"])
        if purpose == "register" and exists:
            raise HTTPException(409, "Ця пошта вже зареєстрована")
        if purpose == "login" and not exists:
            raise HTTPException(404, "Акаунта з такою поштою немає")
        code = otp_issue(data, email, purpose)
        write_db(data)
    sent = send_email(email, "Код підтвердження — Сателіт",
                      f"Ваш код: {code}\nКод дійсний 15 хвилин.")
    resp = {"ok": True, "sent": sent}
    if not sent:
        # Демо-режим (SMTP не налаштовано): код показуємо на сайті.
        resp["demoCode"] = code
    return resp


@app.post("/api/auth/register")
def register(body: RegisterIn):
    email = norm_email(body.email)
    if not EMAIL_RE.match(email):
        raise HTTPException(400, "Вкажіть коректну електронну пошту")
    if len(body.password) < 6:
        raise HTTPException(400, "Пароль має містити щонайменше 6 символів")
    if not body.name.strip():
        raise HTTPException(400, "Вкажіть ім'я")
    with DB_LOCK:
        data = _load()
        if any(u["email"] == email for u in data["users"]):
            raise HTTPException(409, "Ця пошта вже зареєстрована")
        err = otp_check(data, email, "register", body.code.strip())
        if err:
            write_db(data)
            raise HTTPException(400, err)
        user = {
            "id": "u_" + secrets.token_hex(8),
            "email": email,
            "passwordHash": hash_password(body.password),
            "name": body.name.strip(),
            "phone": body.phone.strip(),
            "isAdmin": email == ADMIN_EMAIL,
            "blocked": False,
            "createdAt": now_iso(),
        }
        data["users"].append(user)
        token = make_token(data, user["id"])
        write_db(data)
    return {"token": token, "user": public_user(user)}


@app.post("/api/auth/login")
def login(body: LoginIn):
    email = norm_email(body.email)
    with DB_LOCK:
        data = _load()
        user = next((u for u in data["users"] if u["email"] == email), None)
        if not user or not verify_password(body.password, user["passwordHash"]):
            raise HTTPException(401, "Невірна пошта або пароль")
        if user.get("blocked"):
            raise HTTPException(403, "Акаунт заблоковано. Зверніться до підтримки.")
        if ADMIN_EMAIL and email == ADMIN_EMAIL and not user.get("isAdmin"):
            user["isAdmin"] = True
        if body.code.strip():
            err = otp_check(data, email, "login", body.code.strip())
            if err:
                write_db(data)
                raise HTTPException(400, err)
        token = make_token(data, user["id"])
        write_db(data)
    return {"token": token, "user": public_user(user)}


@app.post("/api/auth/logout")
def logout(request: Request):
    auth = request.headers.get("authorization", "")
    if auth.lower().startswith("bearer "):
        with DB_LOCK:
            data = _load()
            data["tokens"].pop(auth[7:].strip(), None)
            write_db(data)
    return {"ok": True}


@app.get("/api/auth/me")
def me(request: Request):
    data = _load()
    user = require_user(request, data)
    return {"user": public_user(user)}


# ============================================================
# Товари
# ============================================================
def _product_from(body: ProductIn, pid):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Вкажіть назву товару")
    cat = body.category.strip().lower()
    cat = CATEGORY_ALIASES.get(cat, cat)
    if cat not in CATEGORY_IDS:
        cat = "other"
    return {
        "id": pid,
        "name": name,
        "description": body.description.strip(),
        "price": round(float(body.price), 2),
        "category": cat,
        "image": body.image.strip() or DEFAULT_IMAGE,
        "stock": int(body.stock),
        "featured": bool(body.featured),
        "sale": bool(body.sale),
    }


@app.get("/api/products")
def products():
    return {"products": _load()["products"]}


@app.post("/api/admin/products")
def admin_create_product(body: ProductIn, request: Request):
    with DB_LOCK:
        data = _load()
        require_admin(request, data)
        product = _product_from(body, "p-" + secrets.token_hex(4))
        data["products"].append(product)
        write_db(data)
    return {"product": product}


@app.put("/api/admin/products/{pid}")
def admin_update_product(pid: str, body: ProductIn, request: Request):
    with DB_LOCK:
        data = _load()
        require_admin(request, data)
        idx = next((i for i, p in enumerate(data["products"]) if p["id"] == pid), None)
        if idx is None:
            raise HTTPException(404, "Товар не знайдено")
        product = _product_from(body, pid)
        data["products"][idx] = product
        write_db(data)
    return {"product": product}


@app.delete("/api/admin/products/{pid}")
def admin_delete_product(pid: str, request: Request):
    with DB_LOCK:
        data = _load()
        require_admin(request, data)
        before = len(data["products"])
        data["products"] = [p for p in data["products"] if p["id"] != pid]
        if len(data["products"]) == before:
            raise HTTPException(404, "Товар не знайдено")
        write_db(data)
    return {"ok": True}


# ---------- Імпорт товарів із таблиці ----------
HEADER_MAP = {
    "name": {"name", "назва", "найменування", "товар"},
    "description": {"description", "опис"},
    "price": {"price", "ціна", "цена"},
    "category": {"category", "категорія", "категория"},
    "image": {"image", "фото", "зображення", "картинка"},
    "stock": {"stock", "наявність", "залишок", "кількість"},
    "featured": {"featured", "рекомендований", "хіт"},
    "sale": {"sale", "акція"},
}


def _canon_header(h):
    key = str(h or "").strip().lower()
    for canon, names in HEADER_MAP.items():
        if key in names:
            return canon
    return None


def parse_table_bytes(filename, raw):
    """Повертає список рядків: {'_row': номер_рядка_в_файлі, 'name': ..., ...}."""
    fname = (filename or "").lower()
    if fname.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251", errors="replace")   # CSV із Excel на Windows
        delim = ";" if text[:2000].count(";") > text[:2000].count(",") else ","
        table = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    elif fname.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        table = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        raise HTTPException(400, "Підтримуються файли .xlsx або .csv (старий .xls збережіть як .xlsx)")

    if not table:
        return []
    headers = [_canon_header(h) for h in table[0]]
    if "name" not in headers:
        raise HTTPException(400, "У таблиці немає колонки «name» (або «назва»)")
    rows = []
    for idx, cells in enumerate(table[1:], start=2):
        row = {"_row": idx}
        for col, val in zip(headers, cells):
            if col:
                row[col] = val
        rows.append(row)
    return rows


def _to_price(v):
    s = str(v if v is not None else "")
    s = s.replace("₴", "").replace("грн", "").replace(" ", "").replace(",", ".").strip()
    return float(s)


def _to_int(v):
    try:
        return int(float(str(v if v is not None else 0).strip() or 0))
    except ValueError:
        return 0


def _to_bool(v):
    return str(v if v is not None else "").strip().lower() in ("1", "true", "yes", "так", "+", "x")


def _slugify(text):
    slug = re.sub(r"[^\w]+", "-", text.lower()).strip("-")[:40]
    return slug or secrets.token_hex(3)


AI_PROMPT = (
    "Ти перевіряєш товари для інтернет-магазину дрібної електроніки. "
    "Отримуєш JSON-масив рядків таблиці. Для кожного рядка визнач, чи це реальний товар для магазину. "
    "Став ok=false, якщо назва не схожа на товар, ціна явно помилкова (менше 1 або більше 100000 грн), "
    "товар не стосується електроніки чи аксесуарів, або опис містить рекламу сторонніх сайтів чи нецензурщину. "
    "Відповідай ТІЛЬКИ JSON-масивом виду [{\"row\": число, \"ok\": true або false, \"reason\": \"коротко українською\"}]. "
    "Без пояснень і без markdown."
)


def ai_review(items):
    """Повертає (словник row -> (ok, reason), текст помилки або None)."""
    if not OPENROUTER_API_KEY:
        return {}, "ШІ-перевірка вимкнена (не задано OPENROUTER_API_KEY)"
    verdicts = {}
    for i in range(0, len(items), 25):
        chunk = items[i:i + 25]
        try:
            resp = requests.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers={"Authorization": "Bearer " + OPENROUTER_API_KEY,
                         "Content-Type": "application/json"},
                json={
                    "model": OPENROUTER_MODEL,
                    "temperature": 0,
                    "messages": [
                        {"role": "system", "content": AI_PROMPT},
                        {"role": "user", "content": json.dumps(chunk, ensure_ascii=False)},
                    ],
                },
                timeout=90,
            )
            resp.raise_for_status()
            text = resp.json()["choices"][0]["message"]["content"]
            arr = json.loads(text[text.find("["): text.rfind("]") + 1])
            for v in arr:
                verdicts[int(v["row"])] = (bool(v.get("ok", True)), str(v.get("reason", "")))
        except Exception as e:  # мережа, квота, некоректна відповідь
            print("AI review error:", e)
            return verdicts, "ШІ не відповів для частини рядків; ці рядки додано без перевірки"
    return verdicts, None


@app.post("/api/admin/products/import")
def import_products(request: Request, file: UploadFile = File(...), use_ai: int = Query(1)):
    with DB_LOCK:
        require_admin(request, _load())
    raw = file.file.read()
    rows = parse_table_bytes(file.filename, raw)

    candidates, skipped = [], []
    for r in rows:
        values = [v for k, v in r.items() if k != "_row" and v is not None and str(v).strip()]
        if not values:
            continue                                  # порожній рядок — мовчки пропускаємо
        name = str(r.get("name") or "").strip()
        if not name:
            skipped.append({"row": r["_row"], "reason": "немає назви"})
            continue
        try:
            price = _to_price(r.get("price"))
        except ValueError:
            skipped.append({"row": r["_row"], "reason": "некоректна ціна"})
            continue
        if price < 0:
            skipped.append({"row": r["_row"], "reason": "ціна не може бути від'ємною"})
            continue
        cat_raw = str(r.get("category") or "").strip().lower()
        category = CATEGORY_ALIASES.get(cat_raw, cat_raw if cat_raw in CATEGORY_IDS else "other")
        candidates.append({
            "row": r["_row"],
            "name": name,
            "description": str(r.get("description") or "").strip(),
            "price": round(price, 2),
            "category": category,
            "image": str(r.get("image") or "").strip() or DEFAULT_IMAGE,
            "stock": max(_to_int(r.get("stock")), 0),
            "featured": _to_bool(r.get("featured")),
            "sale": _to_bool(r.get("sale")),
        })

    ai_rejected, ai_error = [], None
    if str(use_ai) == "1" and candidates:
        verdicts, ai_error = ai_review(candidates)
        kept = []
        for c in candidates:
            ok, reason = verdicts.get(c["row"], (True, ""))
            if ok:
                kept.append(c)
            else:
                ai_rejected.append({"row": c["row"], "name": c["name"], "reason": reason or "відхилено ШІ"})
        candidates = kept

    added = updated = 0
    with DB_LOCK:
        data = _load()
        for c in candidates:
            pid = "imp-" + _slugify(c["name"])
            product = {k: v for k, v in c.items() if k != "row"}
            product["id"] = pid
            idx = next((i for i, p in enumerate(data["products"]) if p["id"] == pid), None)
            if idx is None:
                data["products"].append(product)
                added += 1
            else:
                data["products"][idx] = product
                updated += 1
        write_db(data)

    resp = {"added": added, "updated": updated, "skipped": skipped, "aiRejected": ai_rejected}
    if ai_error:
        resp["aiError"] = ai_error
    return resp


# ============================================================
# Замовлення
# ============================================================
@app.post("/api/orders")
def create_order(body: OrderIn, request: Request):
    with DB_LOCK:
        data = _load()
        user = user_from_request(request, data)      # гість теж може замовити
        if user and user.get("blocked"):
            raise HTTPException(403, "Акаунт заблоковано. Зверніться до підтримки.")
        if not body.items:
            raise HTTPException(400, "Кошик порожній")
        if body.delivery_method not in DELIVERY_METHODS:
            raise HTTPException(400, "Невірний спосіб доставки")
        name = body.customer_name.strip()
        phone = body.phone.strip()
        if not name:
            raise HTTPException(400, "Вкажіть ім'я")
        if len(re.sub(r"\D", "", phone)) < 9:
            raise HTTPException(400, "Вкажіть коректний номер телефону")

        products_by_id = {p["id"]: p for p in data["products"]}
        lines, total = [], 0.0
        for item in body.items:
            p = products_by_id.get(item.id)
            if not p:
                raise HTTPException(400, f"Товар більше не доступний (id: {item.id})")
            lines.append({"id": p["id"], "name": p["name"], "price": p["price"], "qty": item.qty})
            total += p["price"] * item.qty

        order = {
            "id": "o-" + secrets.token_hex(5),
            "userId": user["id"] if user else None,
            "customerName": name,
            "phone": phone,
            "email": norm_email(body.email),
            "deliveryMethod": body.delivery_method,
            "address": body.address.strip(),
            "items": lines,
            "total": round(total, 2),
            "status": "Нове",
            "createdAt": now_iso(),
        }
        data["orders"].append(order)
        write_db(data)
    return {"order": order}


@app.get("/api/orders")
def my_orders(request: Request):
    data = _load()
    user = require_user(request, data)
    mine = [o for o in data["orders"] if o.get("userId") == user["id"]]
    mine.sort(key=lambda o: o["createdAt"], reverse=True)
    return {"orders": mine}


# ============================================================
# Адмін-панель
# ============================================================
@app.get("/api/admin/orders")
def admin_orders(request: Request):
    data = _load()
    require_admin(request, data)
    orders = sorted(data["orders"], key=lambda o: o["createdAt"], reverse=True)
    return {"orders": orders, "statuses": STATUSES}


@app.patch("/api/admin/orders/{oid}")
def admin_set_order_status(oid: str, body: StatusIn, request: Request):
    if body.status not in STATUSES:
        raise HTTPException(400, "Невідомий статус замовлення")
    with DB_LOCK:
        data = _load()
        require_admin(request, data)
        order = next((o for o in data["orders"] if o["id"] == oid), None)
        if not order:
            raise HTTPException(404, "Замовлення не знайдено")
        order["status"] = body.status
        write_db(data)
    return {"order": order}


@app.get("/api/admin/users")
def admin_users(request: Request):
    data = _load()
    require_admin(request, data)
    return {"users": [public_user(u) for u in data["users"]]}


@app.post("/api/admin/users/{uid}/role")
def admin_set_role(uid: str, body: RoleIn, request: Request):
    with DB_LOCK:
        data = _load()
        me_user = require_admin(request, data)
        target = next((u for u in data["users"] if u["id"] == uid), None)
        if not target:
            raise HTTPException(404, "Користувача не знайдено")
        if target["id"] == me_user["id"] and not body.isAdmin:
            raise HTTPException(400, "Не можна зняти права адміністратора з себе")
        target["isAdmin"] = body.isAdmin
        write_db(data)
    return {"user": public_user(target)}


@app.post("/api/admin/users/{uid}/block")
def admin_set_block(uid: str, body: BlockIn, request: Request):
    with DB_LOCK:
        data = _load()
        me_user = require_admin(request, data)
        target = next((u for u in data["users"] if u["id"] == uid), None)
        if not target:
            raise HTTPException(404, "Користувача не знайдено")
        if target["id"] == me_user["id"] and body.blocked:
            raise HTTPException(400, "Не можна заблокувати себе")
        target["blocked"] = body.blocked
        write_db(data)
    return {"user": public_user(target)}


@app.post("/api/admin/users/{uid}/password")
def admin_set_password(uid: str, body: PasswordIn, request: Request):
    if len(body.password) < 6:
        raise HTTPException(400, "Пароль має містити щонайменше 6 символів")
    with DB_LOCK:
        data = _load()
        require_admin(request, data)
        target = next((u for u in data["users"] if u["id"] == uid), None)
        if not target:
            raise HTTPException(404, "Користувача не знайдено")
        target["passwordHash"] = hash_password(body.password)
        write_db(data)
    return {"ok": True}


@app.delete("/api/admin/users/{uid}")
def admin_delete_user(uid: str, request: Request):
    with DB_LOCK:
        data = _load()
        me_user = require_admin(request, data)
        if uid == me_user["id"]:
            raise HTTPException(400, "Не можна видалити себе")
        if not any(u["id"] == uid for u in data["users"]):
            raise HTTPException(404, "Користувача не знайдено")
        data["users"] = [u for u in data["users"] if u["id"] != uid]
        data["tokens"] = {t: owner for t, owner in data["tokens"].items() if owner != uid}
        write_db(data)
    return {"ok": True}
